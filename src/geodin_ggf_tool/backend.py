from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
import re
from collections import Counter

import openpyxl
from openpyxl.utils.datetime import from_excel


DEFAULT_DATE_COLUMN = "Datum"
DEFAULT_FILTER_COLUMN = "Ionenbilanz"
DEFAULT_FILTER_MIN = -5.0
DEFAULT_FILTER_MAX = 5.0
DEFAULT_DUMMY_DATE = "19000101"
DEFAULT_MAX_SLOTS = 60


@dataclass
class ProcessSettings:
    ggf_input: Path
    xlsx_input: Path
    ggf_output: Path
    report_output: Path | None = None

    date_column: str = DEFAULT_DATE_COLUMN
    filter_column: str = DEFAULT_FILTER_COLUMN
    filter_min: float = DEFAULT_FILTER_MIN
    filter_max: float = DEFAULT_FILTER_MAX

    dummy_date: str = DEFAULT_DUMMY_DATE
    max_slots: int = DEFAULT_MAX_SLOTS

    remove_duplicate_dates: bool = False


@dataclass
class ProcessResult:
    input_ggf: Path
    input_xlsx: Path
    output_ggf: Path
    total_slots_found: int
    valid_excel_dates: int
    final_dates_written: int
    dummy_dates_written: int
    duplicate_dates: list[str]
    file_size_unchanged: bool
    log_lines: list[str]


def normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def get_excel_headers(xlsx_path: Path) -> list[str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.active
    return [normalize_header(cell.value) for cell in ws[1]]


def auto_detect_column(headers: list[str], preferred_name: str) -> str | None:
    """
    Tries to find a column by exact match first, then case-insensitive match.
    """
    if preferred_name in headers:
        return preferred_name

    preferred_lower = preferred_name.lower()
    for header in headers:
        if header.lower() == preferred_lower:
            return header

    return None


def excel_value_to_geodin_date(value) -> str | None:
    """
    Convert Excel/Python/string date values to GeoDIN format YYYYMMDD.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%Y%m%d")

    if isinstance(value, date):
        return value.strftime("%Y%m%d")

    if isinstance(value, (int, float)):
        return from_excel(value).strftime("%Y%m%d")

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y%m%d"):
            try:
                return datetime.strptime(value, fmt).strftime("%Y%m%d")
            except ValueError:
                pass

    raise ValueError(f"Cannot convert Excel date value: {value!r}")


def read_excel_dates(
    xlsx_path: Path,
    date_column: str,
    filter_column: str,
    filter_min: float,
    filter_max: float,
) -> list[str]:
    """
    Read and filter valid dates from the Excel file.
    Keeps only rows where filter_min <= filter_column <= filter_max.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    headers = [normalize_header(cell.value) for cell in ws[1]]

    if date_column not in headers:
        raise ValueError(f"Column '{date_column}' not found. Found: {headers}")

    if filter_column not in headers:
        raise ValueError(f"Column '{filter_column}' not found. Found: {headers}")

    date_col_idx = headers.index(date_column) + 1
    filter_col_idx = headers.index(filter_column) + 1

    dates: list[str] = []

    for row in range(2, ws.max_row + 1):
        date_value = ws.cell(row=row, column=date_col_idx).value
        filter_value = ws.cell(row=row, column=filter_col_idx).value

        if filter_value is None:
            continue

        try:
            filter_value = float(filter_value)
        except (TypeError, ValueError):
            continue

        if filter_min <= filter_value <= filter_max:
            geodin_date = excel_value_to_geodin_date(date_value)
            if geodin_date:
                dates.append(geodin_date)

    return dates


def find_duplicate_dates(dates: list[str]) -> list[str]:
    counts = Counter(dates)
    return sorted([date_value for date_value, count in counts.items() if count > 1])


def remove_duplicates_keep_order(dates: list[str]) -> list[str]:
    seen = set()
    unique_dates = []

    for d in dates:
        if d not in seen:
            unique_dates.append(d)
            seen.add(d)

    return unique_dates


def find_filled_smpdate_slots(data: bytes):
    """
    Find all fixed-length SMPDATE slots in the GGF binary file.

    This intentionally matches only slots where the date already contains
    exactly 8 digits. Therefore the template should have all 60 slots filled,
    usually with real dates or dummy date 19000101.
    """
    pattern = re.compile(
        re.escape("$SMPDATE$='".encode("utf-16-le"))
        + b"((?:[0-9]\x00){8})"
        + re.escape("'".encode("utf-16-le"))
    )

    return list(pattern.finditer(data))


def build_final_date_list(
    excel_dates: list[str],
    max_slots: int,
    dummy_date: str,
) -> list[str]:
    """
    Safe logic:
    1. First fill all positions with dummy date.
    2. Then overwrite the first N positions with valid Excel dates.
    """
    if len(excel_dates) > max_slots:
        raise ValueError(
            f"Excel has {len(excel_dates)} valid dates, "
            f"but the GGF template supports only {max_slots}."
        )

    final_dates = [dummy_date] * max_slots

    for i, geodin_date in enumerate(excel_dates):
        final_dates[i] = geodin_date

    return final_dates


def update_ggf_dates(settings: ProcessSettings) -> ProcessResult:
    """
    Main backend function used by both CLI and GUI.
    """
    settings.ggf_input = Path(settings.ggf_input)
    settings.xlsx_input = Path(settings.xlsx_input)
    settings.ggf_output = Path(settings.ggf_output)

    log_lines: list[str] = []

    excel_dates = read_excel_dates(
        xlsx_path=settings.xlsx_input,
        date_column=settings.date_column,
        filter_column=settings.filter_column,
        filter_min=settings.filter_min,
        filter_max=settings.filter_max,
    )

    duplicate_dates = find_duplicate_dates(excel_dates)

    if settings.remove_duplicate_dates:
        excel_dates = remove_duplicates_keep_order(excel_dates)

    final_dates = build_final_date_list(
        excel_dates=excel_dates,
        max_slots=settings.max_slots,
        dummy_date=settings.dummy_date,
    )

    data = settings.ggf_input.read_bytes()
    output = bytearray(data)

    matches = find_filled_smpdate_slots(data)

    if len(matches) != settings.max_slots:
        raise ValueError(
            f"Expected {settings.max_slots} filled SMPDATE slots, "
            f"found {len(matches)}. "
            f"Please use a template where all SMPDATE slots contain 8-digit dates."
        )

    log_lines.append(f"Input GGF: {settings.ggf_input}")
    log_lines.append(f"Input Excel: {settings.xlsx_input}")
    log_lines.append(f"Output GGF: {settings.ggf_output}")
    log_lines.append("")
    log_lines.append(f"Date column: {settings.date_column}")
    log_lines.append(f"Filter column: {settings.filter_column}")
    log_lines.append(f"Filter range: {settings.filter_min} to {settings.filter_max}")
    log_lines.append("")
    log_lines.append(f"Valid Excel dates after filter: {len(excel_dates)}")
    log_lines.append(f"GGF SMPDATE slots found: {len(matches)}")
    log_lines.append(f"Dummy date used for empty slots: {settings.dummy_date}")
    log_lines.append("")

    if duplicate_dates:
        log_lines.append("Duplicate dates detected:")
        for d in duplicate_dates:
            log_lines.append(f"  - {d}")
        log_lines.append("")

    for i, match in enumerate(matches):
        old_date = match.group(1).decode("utf-16-le")
        new_date = final_dates[i]

        new_block = f"$SMPDATE$='{new_date}'".encode("utf-16-le")

        if len(new_block) != match.end() - match.start():
            raise ValueError(
                f"Length mismatch at slot {i + 1}. "
                f"Old block length: {match.end() - match.start()}, "
                f"new block length: {len(new_block)}. Aborting."
            )

        output[match.start():match.end()] = new_block
        log_lines.append(f"Slot {i + 1:02d}: {old_date} -> {new_date}")

    settings.ggf_output.parent.mkdir(parents=True, exist_ok=True)
    settings.ggf_output.write_bytes(output)

    file_size_unchanged = len(data) == len(output)

    log_lines.append("")
    log_lines.append(f"Saved: {settings.ggf_output}")
    log_lines.append(f"File size unchanged: {file_size_unchanged}")

    if settings.report_output:
        settings.report_output = Path(settings.report_output)
        settings.report_output.parent.mkdir(parents=True, exist_ok=True)
        settings.report_output.write_text("\n".join(log_lines), encoding="utf-8")

    return ProcessResult(
        input_ggf=settings.ggf_input,
        input_xlsx=settings.xlsx_input,
        output_ggf=settings.ggf_output,
        total_slots_found=len(matches),
        valid_excel_dates=len(excel_dates),
        final_dates_written=len(excel_dates),
        dummy_dates_written=settings.max_slots - len(excel_dates),
        duplicate_dates=duplicate_dates,
        file_size_unchanged=file_size_unchanged,
        log_lines=log_lines,
    )


if __name__ == "__main__":
    example_settings = ProcessSettings(
        ggf_input=Path("Br_3_temp_60slots.GGF"),
        xlsx_input=Path("Br_3_chemie_datai.xlsx"),
        ggf_output=Path("Br_3_temp_updated_from_excel.GGF"),
        report_output=Path("Br_3_temp_update_report.txt"),
    )

    result = update_ggf_dates(example_settings)
    print("\n".join(result.log_lines))
